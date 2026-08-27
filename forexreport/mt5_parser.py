"""
Parser exportu historie obchodování z MetaTraderu 5 (soubor ReportHistory-*.xlsx).

Sešit je jednolistový a obsahuje několik za sebou jdoucích sekcí:
hlavička s údaji o účtu, "Pozice", "Pokyny", "Nabídky" (obchody) a "Výsledky".
Parser je odolný vůči české i anglické lokalizaci exportu.
"""

from __future__ import annotations

import re
import warnings
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Dict, List, Optional


# ---------------------------------------------------------------------------
# Datové struktury
# ---------------------------------------------------------------------------


@dataclass
class Position:
    """Jedna uzavřená pozice (kompletní obchod od otevření po uzavření)."""

    ticket: str
    symbol: str
    direction: str                 # 'buy' / 'sell'
    volume: float                  # objem v lotech
    open_time: Optional[datetime]
    close_time: Optional[datetime]
    open_price: Optional[float]
    close_price: Optional[float]
    stop_loss: Optional[float]
    take_profit: Optional[float]
    commission: float
    swap: float
    gross_profit: float            # zisk bez poplatků (sloupec "Zisk")

    @property
    def net_profit(self) -> float:
        """Čistý výsledek obchodu včetně provize a swapu."""
        return self.gross_profit + self.commission + self.swap

    @property
    def duration_minutes(self) -> Optional[float]:
        """Doba držení pozice v minutách; None, pokud chybí některý čas."""
        if not self.open_time or not self.close_time:
            return None
        return (self.close_time - self.open_time).total_seconds() / 60.0

    @property
    def is_win(self) -> bool:
        return self.net_profit > 0

    @property
    def is_loss(self) -> bool:
        return self.net_profit < 0


@dataclass
class Deal:
    """Jednotlivý obchodní příkaz ze sekce "Nabídky" (deals) včetně stavu zůstatku."""

    time: Optional[datetime]
    ticket: str
    symbol: str
    deal_type: str                 # 'buy' / 'sell' / 'balance'
    entry: str                     # 'in' / 'out' / prázdné
    volume: Optional[float]
    price: Optional[float]
    commission: float
    fee: float
    swap: float
    profit: float
    balance: Optional[float]
    comment: str


@dataclass
class AccountInfo:
    """Identifikační údaje účtu z hlavičky exportu."""

    name: str = ""
    account: str = ""
    company: str = ""
    exported_at: str = ""
    currency: str = ""
    server: str = ""


@dataclass
class Mt5Report:
    """Kompletní obsah načteného exportu."""

    account: AccountInfo
    positions: List[Position] = field(default_factory=list)
    deals: List[Deal] = field(default_factory=list)
    summary: Dict[str, Any] = field(default_factory=dict)
    source_file: str = ""

    @property
    def initial_balance(self) -> float:
        """
        Počáteční kapitál účtu.

        Odvozuje se z prvního záznamu typu 'balance' v sekci obchodů; pokud
        takový záznam chybí, použije se zůstatek před prvním obchodem.
        """
        for deal in self.deals:
            if deal.deal_type == "balance":
                return deal.profit
        for deal in self.deals:
            if deal.balance is not None:
                return deal.balance - deal.profit - deal.commission - deal.swap - deal.fee
        return 0.0

    @property
    def deposits(self) -> List[Deal]:
        """Všechny nebchodní pohyby na účtu (vklady, výběry, korekce)."""
        return [d for d in self.deals if d.deal_type == "balance"]


# ---------------------------------------------------------------------------
# Slovníky pro rozpoznání sekcí a sloupců (CZ + EN)
# ---------------------------------------------------------------------------

_SECTION_ALIASES = {
    "positions": {"pozice", "positions"},
    "orders": {"pokyny", "objednávky", "orders"},
    "deals": {"nabídky", "obchody", "deals"},
    "results": {"výsledky", "results"},
}

# Mapa názvu sloupce -> interní klíč. Duplicitní názvy (Čas, Cena) se řeší
# pořadím výskytu v hlavičce.
_COLUMN_ALIASES = {
    "čas": "time",
    "time": "time",
    "čas otevření": "time",
    "open time": "time",
    "pozice": "ticket",
    "position": "ticket",
    "obchod": "ticket",
    "deal": "ticket",
    "pokyn": "order",
    "order": "order",
    "symbol": "symbol",
    "typ": "type",
    "type": "type",
    "směr": "entry",
    "direction": "entry",
    "objem": "volume",
    "volume": "volume",
    "cena": "price",
    "price": "price",
    "s / l": "sl",
    "s/l": "sl",
    "realizovat zisk": "tp",
    "t / p": "tp",
    "t/p": "tp",
    "take profit": "tp",
    "provize": "commission",
    "commission": "commission",
    "poplatek": "fee",
    "fee": "fee",
    "swap": "swap",
    "zisk": "profit",
    "profit": "profit",
    "zůstatek": "balance",
    "balance": "balance",
    "komentář": "comment",
    "comment": "comment",
    "stav": "state",
    "state": "state",
}

_META_ALIASES = {
    "jméno": "name",
    "name": "name",
    "účet": "account",
    "account": "account",
    "firma": "company",
    "company": "company",
    "broker": "company",
    "datum": "exported_at",
    "date": "exported_at",
}


# ---------------------------------------------------------------------------
# Pomocné převodní funkce
# ---------------------------------------------------------------------------


def _norm(value: Any) -> str:
    """Normalizuje buňku na porovnatelný text (bez diakritických rozdílů v mezerách)."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().strip(":").lower()


def _to_float(value: Any, default: float = 0.0) -> float:
    """
    Převede buňku na desetinné číslo.

    Zvládá formáty jako '1 234,56', '0.5 / 0' (částečně plněné pokyny) i prázdné
    buňky; u nečíselných hodnot vrací výchozí hodnotu.
    """
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    # U pokynů má objem tvar "požadováno / plněno" – bereme první číslo
    if "/" in text:
        text = text.split("/")[0]
    text = text.replace(" ", "").replace(" ", "").replace(",", ".")
    text = re.sub(r"[^0-9.\-+eE]", "", text)
    try:
        return float(text)
    except ValueError:
        return default


def _to_optional_float(value: Any) -> Optional[float]:
    """Jako _to_float, ale prázdné či nečíselné hodnoty vrací jako None."""
    if value is None or str(value).strip() == "":
        return None
    result = _to_float(value, default=float("nan"))
    return None if result != result else result


_DATE_FORMATS = (
    "%Y.%m.%d %H:%M:%S",
    "%Y.%m.%d %H:%M",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%d.%m.%Y %H:%M:%S",
    "%Y.%m.%d",
    "%Y-%m-%d",
)


def _to_datetime(value: Any) -> Optional[datetime]:
    """Převede buňku na datetime; podporuje běžné formáty MT5 i nativní datum z Excelu."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _header_map(row: List[Any]) -> Dict[str, int]:
    """
    Vytvoří mapu interní klíč -> index sloupce z hlavičkového řádku.

    Duplicitní názvy dostávají příponu '_2' (v sekci pozic jsou dva sloupce
    "Čas" a dva "Cena" – otevírací a zavírací).
    """
    mapping: Dict[str, int] = {}
    seen: Dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _COLUMN_ALIASES.get(_norm(cell))
        if not key:
            continue
        seen[key] = seen.get(key, 0) + 1
        name = key if seen[key] == 1 else f"{key}_{seen[key]}"
        mapping[name] = index
    return mapping


def _cell(row: List[Any], mapping: Dict[str, int], key: str) -> Any:
    """Bezpečně vytáhne hodnotu sloupce podle interního klíče."""
    index = mapping.get(key)
    if index is None or index >= len(row):
        return None
    return row[index]


def _is_blank(row: List[Any]) -> bool:
    """Řádek je prázdný, pokud neobsahuje žádnou neprázdnou buňku."""
    return all(cell is None or str(cell).strip() == "" for cell in row)


# ---------------------------------------------------------------------------
# Hlavní načtení
# ---------------------------------------------------------------------------


def load_report(path: str) -> Mt5Report:
    """
    Načte MT5 export a vrátí strukturovaná data.

    :param path: cesta k souboru ReportHistory-*.xlsx
    :raises ValueError: pokud sešit neobsahuje rozpoznatelnou sekci pozic ani obchodů
    """
    import openpyxl

    # openpyxl varuje u MT5 sešitů na chybějící výchozí styl – pro uživatele nepodstatné
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        raise ValueError(
            f"Soubor {path} se nepodařilo otevřít jako sešit Excelu. "
            "Očekává se export 'Historie obchodování' z MetaTraderu 5 ve formátu .xlsx. "
            f"Podrobnost: {exc}"
        ) from exc

    sheet = workbook.active
    rows: List[List[Any]] = [
        [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        for r in range(1, sheet.max_row + 1)
    ]

    account = _parse_account(rows)
    sections = _locate_sections(rows)

    positions = _parse_positions(rows, sections.get("positions"))
    deals = _parse_deals(rows, sections.get("deals"))
    summary = _parse_summary(rows, sections.get("results"))

    if not positions and not deals:
        raise ValueError(
            f"V souboru {path} se nepodařilo najít sekci s pozicemi ani obchody. "
            "Očekává se export 'Historie obchodování' z MetaTraderu 5."
        )

    # Pokud sekce pozic chybí (některé exporty ji nemají), sestavíme pozice z obchodů
    if not positions and deals:
        positions = _positions_from_deals(deals)

    report = Mt5Report(
        account=account,
        positions=sorted(positions, key=lambda p: p.close_time or datetime.min),
        deals=deals,
        summary=summary,
        source_file=path,
    )
    return report


def _parse_account(rows: List[List[Any]]) -> AccountInfo:
    """Vytáhne údaje o účtu z prvních řádků sešitu (hledá dvojice popisek/hodnota)."""
    info = AccountInfo()
    for row in rows[:12]:
        label = _norm(row[0] if row else None)
        key = _META_ALIASES.get(label)
        if not key:
            continue
        # Hodnota je v první neprázdné buňce vpravo od popisku (bývá sloučená)
        value = next((str(c).strip() for c in row[1:] if c is not None and str(c).strip()), "")
        setattr(info, key, value)

    # Z řetězce účtu typu "498754 (USD, FTTrading-Server, demo, Hedge)" vytáhneme měnu a server
    match = re.match(r"^\s*(\S+)\s*\((.*)\)\s*$", info.account)
    if match:
        parts = [p.strip() for p in match.group(2).split(",")]
        if parts:
            info.currency = parts[0]
        if len(parts) > 1:
            info.server = parts[1]
    return info


def _locate_sections(rows: List[List[Any]]) -> Dict[str, int]:
    """
    Najde indexy řádků, na kterých začínají jednotlivé sekce.

    Sekční řádek poznáme podle toho, že má vyplněnou pouze první buňku a její
    text odpovídá známému názvu sekce.
    """
    found: Dict[str, int] = {}
    for index, row in enumerate(rows):
        if not row:
            continue
        label = _norm(row[0])
        if not label:
            continue
        rest_empty = all(c is None or str(c).strip() == "" for c in row[1:])
        if not rest_empty:
            continue
        for section, aliases in _SECTION_ALIASES.items():
            if label in aliases and section not in found:
                found[section] = index
    return found


def _section_rows(rows: List[List[Any]], start: Optional[int]):
    """
    Generátor datových řádků sekce.

    Vrací dvojici (mapa sloupců, řádek). Sekce končí prvním prázdným řádkem
    nebo řádkem, který vypadá jako začátek další sekce.
    """
    if start is None:
        return
    if start + 1 >= len(rows):
        return
    mapping = _header_map(rows[start + 1])
    if not mapping:
        return
    for row in rows[start + 2:]:
        if _is_blank(row):
            break
        label = _norm(row[0])
        if any(label in aliases for aliases in _SECTION_ALIASES.values()):
            break
        yield mapping, row


def _parse_positions(rows: List[List[Any]], start: Optional[int]) -> List[Position]:
    """Načte sekci "Pozice" – uzavřené obchody s otevíracími a zavíracími údaji."""
    positions: List[Position] = []
    for mapping, row in _section_rows(rows, start):
        ticket = _cell(row, mapping, "ticket")
        direction = _norm(_cell(row, mapping, "type"))
        # Souhrnné řádky na konci sekce nemají ticket ani směr – ty přeskočíme
        if not ticket or direction not in ("buy", "sell"):
            continue
        positions.append(
            Position(
                ticket=str(ticket).strip(),
                symbol=str(_cell(row, mapping, "symbol") or "").strip(),
                direction=direction,
                volume=_to_float(_cell(row, mapping, "volume")),
                open_time=_to_datetime(_cell(row, mapping, "time")),
                close_time=_to_datetime(_cell(row, mapping, "time_2")),
                open_price=_to_optional_float(_cell(row, mapping, "price")),
                close_price=_to_optional_float(_cell(row, mapping, "price_2")),
                stop_loss=_to_optional_float(_cell(row, mapping, "sl")),
                take_profit=_to_optional_float(_cell(row, mapping, "tp")),
                commission=_to_float(_cell(row, mapping, "commission")),
                swap=_to_float(_cell(row, mapping, "swap")),
                gross_profit=_to_float(_cell(row, mapping, "profit")),
            )
        )
    return positions


def _parse_deals(rows: List[List[Any]], start: Optional[int]) -> List[Deal]:
    """Načte sekci obchodů ("Nabídky"/"Deals") včetně vkladů a průběžného zůstatku."""
    deals: List[Deal] = []
    for mapping, row in _section_rows(rows, start):
        ticket = _cell(row, mapping, "ticket")
        time = _to_datetime(_cell(row, mapping, "time"))
        # Poslední řádek sekce je součet – nemá čas ani ticket
        if not ticket or time is None:
            continue
        deals.append(
            Deal(
                time=time,
                ticket=str(ticket).strip(),
                symbol=str(_cell(row, mapping, "symbol") or "").strip(),
                deal_type=_norm(_cell(row, mapping, "type")),
                entry=_norm(_cell(row, mapping, "entry")),
                volume=_to_optional_float(_cell(row, mapping, "volume")),
                price=_to_optional_float(_cell(row, mapping, "price")),
                commission=_to_float(_cell(row, mapping, "commission")),
                fee=_to_float(_cell(row, mapping, "fee")),
                swap=_to_float(_cell(row, mapping, "swap")),
                profit=_to_float(_cell(row, mapping, "profit")),
                balance=_to_optional_float(_cell(row, mapping, "balance")),
                comment=str(_cell(row, mapping, "comment") or "").strip(),
            )
        )
    return deals


def _parse_summary(rows: List[List[Any]], start: Optional[int]) -> Dict[str, Any]:
    """
    Načte sekci "Výsledky" jako volný slovník popisek -> hodnota.

    Řádky obsahují více dvojic vedle sebe, proto procházíme buňky po dvojicích
    popisek/první neprázdná hodnota vpravo.
    """
    summary: Dict[str, Any] = {}
    if start is None:
        return summary
    for row in rows[start + 1:]:
        if _is_blank(row):
            continue
        index = 0
        while index < len(row):
            cell = row[index]
            text = str(cell).strip() if cell is not None else ""
            if text.endswith(":"):
                label = text.rstrip(":").strip()
                value = None
                for probe in range(index + 1, len(row)):
                    if row[probe] is not None and str(row[probe]).strip():
                        # Další popisek znamená, že hodnota chybí
                        if str(row[probe]).strip().endswith(":"):
                            break
                        value = row[probe]
                        index = probe
                        break
                summary[label] = value
            index += 1
    return summary


def _positions_from_deals(deals: List[Deal]) -> List[Position]:
    """
    Nouzová rekonstrukce pozic z obchodů, pokud export sekci pozic neobsahuje.

    Páruje vstupní ('in') a výstupní ('out') obchod podle symbolu v pořadí FIFO.
    """
    open_deals: Dict[str, List[Deal]] = {}
    positions: List[Position] = []
    for deal in deals:
        if deal.deal_type not in ("buy", "sell"):
            continue
        if deal.entry == "in":
            open_deals.setdefault(deal.symbol, []).append(deal)
        elif deal.entry == "out":
            queue = open_deals.get(deal.symbol)
            if not queue:
                continue
            entry_deal = queue.pop(0)
            positions.append(
                Position(
                    ticket=entry_deal.ticket,
                    symbol=deal.symbol,
                    # Směr pozice je opačný než uzavírací obchod
                    direction="buy" if deal.deal_type == "sell" else "sell",
                    volume=deal.volume or 0.0,
                    open_time=entry_deal.time,
                    close_time=deal.time,
                    open_price=entry_deal.price,
                    close_price=deal.price,
                    stop_loss=None,
                    take_profit=None,
                    commission=entry_deal.commission + deal.commission,
                    swap=entry_deal.swap + deal.swap,
                    gross_profit=deal.profit,
                )
            )
    return positions
